"""画像内の図形検出と PPTX への変換ロジック。"""

from __future__ import annotations

import io
import math
from dataclasses import dataclass, field
from typing import List, Optional, Tuple

import cv2
import numpy as np
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.util import Emu


# OpenCV の近似頂点数 -> PowerPoint オートシェイプ種別の対応
SHAPE_BY_VERTICES = {
    3: MSO_SHAPE.ISOSCELES_TRIANGLE,
    4: MSO_SHAPE.RECTANGLE,
    5: MSO_SHAPE.REGULAR_PENTAGON,
    6: MSO_SHAPE.HEXAGON,
}


@dataclass
class DetectedShape:
    """画像内で検出された 1 つの図形（ピクセル座標系）。"""

    kind: str                       # "rectangle" | "ellipse" | "triangle" | ...
    bbox: Tuple[int, int, int, int]  # x, y, w, h（ピクセル）
    fill: Tuple[int, int, int]      # RGB
    line: Tuple[int, int, int]      # RGB（輪郭色）
    vertices: int = 0

    @property
    def area(self) -> int:
        _, _, w, h = self.bbox
        return w * h


def _dominant_color(image: np.ndarray, mask: np.ndarray) -> Tuple[int, int, int]:
    """マスク領域の平均色を RGB で返す（image は BGR）。"""
    pixels = image[mask.astype(bool)]
    if pixels.size == 0:
        return (255, 255, 255)
    b, g, r = pixels.mean(axis=0)
    return (int(r), int(g), int(b))


def detect_shapes_in_image(
    image_bytes: bytes,
    *,
    min_area_ratio: float = 0.002,
    max_area_ratio: float = 0.98,
    approx_epsilon: float = 0.03,
) -> List[DetectedShape]:
    """画像バイト列から単純図形を検出する。

    Args:
        image_bytes: PNG/JPEG などの画像バイト列。
        min_area_ratio: 画像全体に対する図形の最小面積比（ノイズ除去）。
        max_area_ratio: 最大面積比（画像枠そのものを拾わないため）。
        approx_epsilon: 輪郭近似の許容誤差（周長に対する比）。

    Returns:
        検出した DetectedShape のリスト（面積降順）。
    """
    arr = np.frombuffer(image_bytes, dtype=np.uint8)
    image = cv2.imdecode(arr, cv2.IMREAD_COLOR)
    if image is None:
        return []

    h, w = image.shape[:2]
    total_area = float(h * w)
    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    blurred = cv2.GaussianBlur(gray, (5, 5), 0)
    edges = cv2.Canny(blurred, 50, 150)
    # 図形の輪郭を閉じるために膨張
    edges = cv2.dilate(edges, np.ones((3, 3), np.uint8), iterations=1)

    contours, _ = cv2.findContours(edges, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)

    results: List[DetectedShape] = []
    for cnt in contours:
        area = cv2.contourArea(cnt)
        if area < total_area * min_area_ratio or area > total_area * max_area_ratio:
            continue

        peri = cv2.arcLength(cnt, True)
        if peri == 0:
            continue
        approx = cv2.approxPolyDP(cnt, approx_epsilon * peri, True)
        verts = len(approx)
        x, y, bw, bh = cv2.boundingRect(approx)

        # 円形度で楕円/円を判定
        circularity = 4 * math.pi * area / (peri * peri)
        if verts > 6 and circularity > 0.7:
            kind = "ellipse"
        elif verts == 4:
            kind = "rectangle"
        elif verts == 3:
            kind = "triangle"
        elif verts in SHAPE_BY_VERTICES:
            kind = {5: "pentagon", 6: "hexagon"}[verts]
        else:
            # 不明な多角形はバウンディングボックス矩形として扱う
            kind = "rectangle"
            verts = 4

        mask = np.zeros(gray.shape, dtype=np.uint8)
        cv2.drawContours(mask, [cnt], -1, 255, thickness=-1)
        fill = _dominant_color(image, mask)

        outline_mask = np.zeros(gray.shape, dtype=np.uint8)
        cv2.drawContours(outline_mask, [cnt], -1, 255, thickness=2)
        line = _dominant_color(image, outline_mask)

        results.append(
            DetectedShape(
                kind=kind,
                bbox=(x, y, bw, bh),
                fill=fill,
                line=line,
                vertices=verts,
            )
        )

    results.sort(key=lambda s: s.area, reverse=True)
    return results


_PPTX_SHAPE_ENUM = {
    "rectangle": MSO_SHAPE.RECTANGLE,
    "ellipse": MSO_SHAPE.OVAL,
    "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE,
    "pentagon": MSO_SHAPE.REGULAR_PENTAGON,
    "hexagon": MSO_SHAPE.HEXAGON,
}


@dataclass
class ConversionReport:
    """変換結果のサマリ。"""

    pictures_processed: int = 0
    shapes_created: int = 0
    per_slide: List[int] = field(default_factory=list)


def _add_shape_for_detection(
    slide_shapes,
    det: DetectedShape,
    pic_left: int,
    pic_top: int,
    pic_width: int,
    pic_height: int,
    img_w: int,
    img_h: int,
) -> None:
    """検出図形を、元画像の位置・サイズに合わせて EMU 座標へスケールして追加する。"""
    x, y, bw, bh = det.bbox
    sx = pic_width / img_w
    sy = pic_height / img_h

    left = Emu(int(pic_left + x * sx))
    top = Emu(int(pic_top + y * sy))
    width = Emu(int(bw * sx))
    height = Emu(int(bh * sy))

    enum = _PPTX_SHAPE_ENUM.get(det.kind, MSO_SHAPE.RECTANGLE)
    shape = slide_shapes.add_shape(enum, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*det.fill)
    shape.line.color.rgb = RGBColor(*det.line)
    shape.name = f"Vectorized-{det.kind}"


def representative_color(image_bgr: np.ndarray, bbox: Tuple[int, int, int, int]) -> Tuple[int, int, int]:
    """画像の指定矩形領域の代表色（中央値）を RGB で返す。

    平均ではなく中央値を使うことで、縁のアンチエイリアスやノイズの影響を抑え、
    領域内で支配的な色を拾いやすくする。image_bgr は OpenCV の BGR 画像。
    """
    x, y, w, h = bbox
    h_img, w_img = image_bgr.shape[:2]
    x0 = max(0, min(x, w_img - 1))
    y0 = max(0, min(y, h_img - 1))
    x1 = max(x0 + 1, min(x + w, w_img))
    y1 = max(y0 + 1, min(y + h, h_img))
    patch = image_bgr[y0:y1, x0:x1].reshape(-1, 3)
    if patch.size == 0:
        return (255, 255, 255)
    b, g, r = np.median(patch, axis=0)
    return (int(r), int(g), int(b))


def add_editable_rectangle(
    slide_shapes,
    region_px: Tuple[int, int, int, int],
    pic_left: int,
    pic_top: int,
    pic_width: int,
    pic_height: int,
    img_w: int,
    img_h: int,
    fill_rgb: Tuple[int, int, int],
    *,
    line_rgb: Optional[Tuple[int, int, int]] = None,
    name: str = "Region",
):
    """ユーザーが選択した画像内領域を、編集可能な矩形オートシェイプとして追加する。

    region_px は画像ピクセル座標系の (x, y, w, h)。画像上の位置を、スライド上での
    画像の表示位置・サイズ（EMU）に合わせてスケールして配置する。

    Returns:
        追加した shape オブジェクト。
    """
    x, y, w, h = region_px
    sx = pic_width / img_w
    sy = pic_height / img_h

    left = Emu(int(pic_left + x * sx))
    top = Emu(int(pic_top + y * sy))
    width = Emu(max(1, int(w * sx)))
    height = Emu(max(1, int(h * sy)))

    shape = slide_shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*fill_rgb)
    if line_rgb is None:
        shape.line.fill.background()  # 枠線なし
    else:
        shape.line.color.rgb = RGBColor(*line_rgb)
    shape.name = name
    return shape


def convert_presentation(
    input_path: str,
    output_path: str,
    *,
    remove_original_picture: bool = False,
    progress=None,
    **detect_kwargs,
) -> ConversionReport:
    """PPTX を読み込み、各画像内の図形をオートシェイプに変換して保存する。

    Args:
        input_path: 入力 .pptx パス。
        output_path: 出力 .pptx パス。
        remove_original_picture: True なら元の画像を削除する（既定は残す）。
        progress: callable(done, total, message) 形式の進捗コールバック（任意）。
        **detect_kwargs: detect_shapes_in_image へ渡す検出パラメータ。

    Returns:
        ConversionReport。
    """
    prs = Presentation(input_path)
    report = ConversionReport()

    slides = list(prs.slides)
    total = len(slides)
    for idx, slide in enumerate(slides):
        created_in_slide = 0
        # イテレート中に追加・削除するのでリスト化
        pictures = [s for s in slide.shapes if s.shape_type == 13]  # 13 = PICTURE
        for pic in pictures:
            try:
                image_bytes = pic.image.blob
            except Exception:
                continue
            shapes = detect_shapes_in_image(image_bytes, **detect_kwargs)
            if not shapes:
                continue

            arr = np.frombuffer(image_bytes, dtype=np.uint8)
            img = cv2.imdecode(arr, cv2.IMREAD_COLOR)
            if img is None:
                continue
            img_h, img_w = img.shape[:2]

            for det in shapes:
                _add_shape_for_detection(
                    slide.shapes, det,
                    pic.left, pic.top, pic.width, pic.height,
                    img_w, img_h,
                )
                created_in_slide += 1

            report.pictures_processed += 1

            if remove_original_picture:
                pic._element.getparent().remove(pic._element)

        report.shapes_created += created_in_slide
        report.per_slide.append(created_in_slide)
        if progress:
            progress(idx + 1, total, f"スライド {idx + 1}/{total}: 図形 {created_in_slide} 個")

    prs.save(output_path)
    return report
