"""図形検出と PPTX 変換のテスト。合成画像/PPTX を生成して検証する。"""

import io

import cv2
import numpy as np
import pytest
from pptx import Presentation
from pptx.util import Inches

from pptx import Presentation as _P
from pptx.util import Inches as _In

from pptx_vectorizer.core import (
    convert_presentation,
    detect_shapes_in_image,
    add_editable_rectangle,
    representative_color,
    detect_shapes_in_region,
)


def _make_image_with_shapes() -> bytes:
    """白背景に矩形と円を描いた PNG バイト列を返す。"""
    img = np.full((400, 400, 3), 255, dtype=np.uint8)
    cv2.rectangle(img, (40, 40), (180, 180), (0, 0, 255), -1)   # 赤い矩形
    cv2.circle(img, (300, 300), 70, (255, 0, 0), -1)            # 青い円
    ok, buf = cv2.imencode(".png", img)
    assert ok
    return buf.tobytes()


def test_detect_shapes_finds_rectangle_and_ellipse():
    shapes = detect_shapes_in_image(_make_image_with_shapes())
    kinds = {s.kind for s in shapes}
    assert "rectangle" in kinds
    assert "ellipse" in kinds


def test_detect_empty_for_blank_image():
    img = np.full((200, 200, 3), 255, dtype=np.uint8)
    _, buf = cv2.imencode(".png", img)
    assert detect_shapes_in_image(buf.tobytes()) == []


def test_detect_invalid_bytes_returns_empty():
    assert detect_shapes_in_image(b"not an image") == []


def _make_pptx_with_picture(path: str) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白レイアウト
    slide.shapes.add_picture(
        io.BytesIO(_make_image_with_shapes()),
        Inches(1), Inches(1), Inches(4), Inches(4),
    )
    prs.save(path)


def test_convert_presentation_creates_shapes(tmp_path):
    src = tmp_path / "in.pptx"
    dst = tmp_path / "out.pptx"
    _make_pptx_with_picture(str(src))

    report = convert_presentation(str(src), str(dst))

    assert report.pictures_processed == 1
    assert report.shapes_created >= 2
    assert dst.exists()

    # 出力にオートシェイプが含まれることを確認
    out = Presentation(str(dst))
    names = [s.name for s in out.slides[0].shapes]
    assert any(n.startswith("Vectorized-") for n in names)


def test_representative_color_picks_region_color():
    img = np.full((100, 100, 3), 255, dtype=np.uint8)
    img[10:40, 10:40] = (0, 0, 255)  # BGR 赤
    r, g, b = representative_color(img, (10, 10, 30, 30))
    assert (r, g, b) == (255, 0, 0)  # RGB 赤


def test_representative_color_clamps_out_of_bounds():
    img = np.full((50, 50, 3), 128, dtype=np.uint8)
    # 画像外にはみ出す領域でも例外なく代表色を返す
    assert representative_color(img, (40, 40, 100, 100)) == (128, 128, 128)


def test_add_editable_rectangle_places_scaled_shape():
    prs = _P()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    img = np.zeros((200, 400, 3), dtype=np.uint8)  # img_w=400, img_h=200
    pic = slide.shapes.add_picture(
        io.BytesIO(_make_image_with_shapes()), _In(1), _In(1), _In(4), _In(2),
    )
    before = len(slide.shapes._spTree)
    shape = add_editable_rectangle(
        slide.shapes, (100, 50, 200, 100),
        pic.left, pic.top, pic.width, pic.height,
        400, 200, (10, 20, 30), name="Region-1",
    )
    # 画像の (100,50) は画像幅の 1/4, 高さの 1/4 -> 表示位置も比例
    assert shape.left == pic.left + int(100 * pic.width / 400)
    assert shape.top == pic.top + int(50 * pic.height / 200)
    assert shape.width == int(200 * pic.width / 400)
    assert shape.name == "Region-1"
    assert len(slide.shapes._spTree) == before + 1


def test_detect_shapes_in_region_finds_colored_shape_not_white():
    # 白背景に赤い矩形。背景が多めの範囲を選んでも、白でなく赤い図形を検出する
    img = np.full((300, 400, 3), 255, dtype=np.uint8)
    cv2.rectangle(img, (40, 40), (160, 160), (0, 0, 255), -1)  # BGR 赤
    dets = detect_shapes_in_region(img, (20, 20, 180, 180))
    assert dets, "範囲内の図形が検出されるべき"
    d = dets[0]
    assert d.kind == "rectangle"
    r, g, b = d.fill
    assert r > 200 and g < 80 and b < 80  # 赤（白ではない）
    # 座標は画像全体系に戻っている（オフセット適用済み）
    bx, by, bw, bh = d.bbox
    assert bx >= 20 and by >= 20


def test_detect_shapes_in_region_empty_on_blank():
    img = np.full((200, 200, 3), 255, dtype=np.uint8)
    assert detect_shapes_in_region(img, (10, 10, 100, 100)) == []


def test_convert_remove_original(tmp_path):
    src = tmp_path / "in.pptx"
    dst = tmp_path / "out.pptx"
    _make_pptx_with_picture(str(src))

    convert_presentation(str(src), str(dst), remove_original_picture=True)

    out = Presentation(str(dst))
    pic_count = sum(1 for s in out.slides[0].shapes if s.shape_type == 13)
    assert pic_count == 0
